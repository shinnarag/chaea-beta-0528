from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "ChaeA_Complete_Identity_Dossier.md"
OUTPUT = ROOT / "docs" / "ChaeA_Complete_Identity_Dossier.docx"

BLUE = RGBColor(0x34, 0x78, 0xA5)
DARK = RGBColor(0x20, 0x23, 0x26)
MUTED = RGBColor(0x6D, 0x74, 0x7B)
PEACH = RGBColor(0xD8, 0x5F, 0x73)
HEADER_FILL = "EAF4F8"
SOFT_FILL = "FFF8F1"
BORDER = "D8E4EA"


def set_run_font(run, size: float | None = None, bold: bool | None = None, color=None):
    run.font.name = "Arial"
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color

    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), "Arial")
    r_fonts.set(qn("w:hAnsi"), "Arial")
    r_fonts.set(qn("w:eastAsia"), "Malgun Gothic")


def set_style_font(style, size: float, color=None, bold=False):
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color

    r_pr = style._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), "Arial")
    r_fonts.set(qn("w:hAnsi"), "Arial")
    r_fonts.set(qn("w:eastAsia"), "Malgun Gothic")


def set_spacing(paragraph, before=0, after=6, line=1.25):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_inline(paragraph, text: str, *, size: float | None = None):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text.strip())
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True, color=DARK)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=size, color=BLUE)
            run.font.name = "Courier New"
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, color=DARK)


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = True
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = qn(f"w:{edge}")
        elem = borders.find(tag)
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "4")
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), BORDER)


def configure(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.right_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)

    styles = doc.styles
    set_style_font(styles["Normal"], 10.5, DARK)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.25

    set_style_font(styles["Title"], 24, DARK, True)
    styles["Title"].paragraph_format.space_after = Pt(8)

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 11.5, PEACH, 9, 4),
    ]:
        style = styles[name]
        set_style_font(style, size, color, True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.18

    for list_style in ["List Bullet", "List Number"]:
        style = styles[list_style]
        set_style_font(style, 10.5, DARK)
        style.paragraph_format.left_indent = Inches(0.35)
        style.paragraph_format.first_line_indent = Inches(-0.18)
        style.paragraph_format.space_after = Pt(4)


def parse_table(lines: list[str], start: int):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        row = [cell.strip() for cell in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in row):
            rows.append(row)
        i += 1
    return rows, i


def add_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return

    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    style_table(table)
    for r_idx, row in enumerate(rows):
        for c_idx, cell_text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            shade_cell(cell, HEADER_FILL if r_idx == 0 else SOFT_FILL if c_idx == 0 else "FFFFFF")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_spacing(paragraph, after=2, line=1.15)
            add_inline(paragraph, cell_text, size=10)
            for run in paragraph.runs:
                if r_idx == 0 or c_idx == 0:
                    run.bold = True
                if r_idx == 0:
                    run.font.color.rgb = BLUE

    doc.add_paragraph()


def add_cover(doc: Document, lines: list[str]):
    title = doc.add_paragraph(style="Title")
    add_inline(title, "ChaeA Complete Identity Dossier")

    subtitle = doc.add_paragraph()
    set_spacing(subtitle, after=10)
    run = subtitle.add_run("채아 전체 정보 및 정체성 통합 문서")
    set_run_font(run, size=12.5, bold=True, color=BLUE)

    for line in lines[1:5]:
        if not line.strip():
            continue
        p = doc.add_paragraph()
        set_spacing(p, after=2)
        add_inline(p, line.strip(), size=10.5)

    divider = doc.add_paragraph()
    set_spacing(divider, before=8, after=12)
    run = divider.add_run("LINE / Room Cover / Lyric Diary / Word Collection")
    set_run_font(run, size=9.5, bold=True, color=PEACH)


def build_docx():
    markdown = SOURCE.read_text(encoding="utf-8")
    lines = markdown.splitlines()

    doc = Document()
    configure(doc)
    add_cover(doc, lines)

    first_rule = next((idx for idx, line in enumerate(lines) if line.strip() == "---"), 0)
    i = first_rule + 1
    while i < len(lines):
        line = lines[i].strip()

        if not line or line == "---":
            i += 1
            continue

        if line.startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue

        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            style = "Heading 1" if level == 1 else "Heading 2" if level == 2 else "Heading 3"
            p = doc.add_paragraph(style=style)
            add_inline(p, heading.group(2))
            i += 1
            continue

        bullet = re.match(r"^-\s+(.+)$", line)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, bullet.group(1))
            i += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.+)$", line)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, numbered.group(1))
            i += 1
            continue

        p = doc.add_paragraph()
        set_spacing(p)
        add_inline(p, line)
        i += 1

    doc.core_properties.title = "ChaeA Complete Identity Dossier"
    doc.core_properties.subject = "채아 정체성 및 챗봇 기준 문서"
    doc.core_properties.author = "Sound Republica / Codex"
    doc.core_properties.keywords = "ChaeA, 윤채아, YOON CHAEA, virtual artist"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_docx())
