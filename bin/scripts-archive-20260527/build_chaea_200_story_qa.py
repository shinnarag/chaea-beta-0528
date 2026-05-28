from __future__ import annotations

from datetime import date
from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUT_DOCX = DOCS / "ChaeA_200_Story_Linked_QA.docx"
OUT_MD = DOCS / "ChaeA_200_Story_Linked_QA.md"

sys.path.insert(0, str(ROOT / "scripts"))
from build_chaea_life_story_archive import story_rows  # noqa: E402

INK = RGBColor(0x20, 0x23, 0x26)
MUTED = RGBColor(0x6D, 0x74, 0x7B)
BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
HEADER_FILL = "E8EEF5"
BORDER = "DADCE0"


def rfonts(run, latin="Calibri", east_asia="Malgun Gothic"):
    run.font.name = latin
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), east_asia)


def style_run(run, size=None, color=None, bold=None, italic=None):
    rfonts(run)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.85)
    section.header_distance = Inches(0.42)
    section.footer_distance = Inches(0.42)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Malgun Gothic")
    normal.font.size = Pt(10)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.18

    title = styles["Title"]
    title.font.name = "Calibri"
    title._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Malgun Gothic")
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 14, 8),
        ("Heading 2", 13, BLUE, 10, 5),
        ("Heading 3", 11.5, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Malgun Gothic")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15


def cell_margins(cell, top=70, start=100, bottom=70, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, val in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(table, color=BORDER):
    tbl_pr = table._tbl.tblPr
    b = tbl_pr.first_child_found_in("w:tblBorders")
    if b is None:
        b = OxmlElement("w:tblBorders")
        tbl_pr.append(b)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = b.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            b.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def grid(table, widths, indent_dxa=120):
    total = sum(int(w * 1440) for w in widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_grid = table._tbl.tblGrid
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width * 1440)))
        tbl_grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Inches(widths[i])
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(int(widths[i] * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths, font_size=8.8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    borders(table)
    grid(table, widths)
    for i, header in enumerate(headers):
        c = table.rows[0].cells[i]
        c.text = ""
        shade(c, HEADER_FILL)
        cell_margins(c)
        r = c.paragraphs[0].add_run(header)
        style_run(r, font_size, DARK_BLUE, True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            cell_margins(cells[i])
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.12
            r = p.add_run(value)
            style_run(r, font_size, INK)
    grid(table, widths)
    doc.add_paragraph()


def para(doc, text, style=None, size=None, color=None, bold=False):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    style_run(r, size, color or INK, bold)


def split_sentences(text: str):
    clean = re.sub(r"\s+", " ", text.strip())
    parts = re.split(r"(?<=[.!?。요다죠까])\s+", clean)
    return [p.strip() for p in parts if p.strip()]


def compact(text: str, limit=210):
    text = re.sub(r"\s+", " ", text.strip())
    if len(text) <= limit:
        return text
    cut = text[:limit].rsplit(" ", 1)[0]
    return cut + "..."


def story_answer_a(body):
    first = body.split("\n\n")[0]
    return compact(first, 260)


def story_answer_b(body, data):
    sentences = split_sentences(body)
    tail = " ".join(sentences[-3:]) if len(sentences) >= 3 else body
    return compact(f"그 경험은 {data} 쪽으로 남아요. {tail}", 260)


def story_answer_c(title, data):
    return compact(
        f"이 이야기는 채아가 지금 어떤 취향과 태도로 말하고 노래하는지 설명하는 배경이에요. "
        f"'{title}'의 기억은 {data}와 연결되고, 이후 LINE, 서울 원룸, 커버 아카이브 안에서 다시 나타나요.",
        260,
    )


def build_story_questions():
    rows = []
    q_no = 1
    for no, title, age, place, body, data in story_rows():
        rows.append({
            "no": q_no,
            "priority": "A",
            "story": f"{no}. {title}",
            "question": f"채아에게 '{title}'은 어떤 기억이야?",
            "answer": story_answer_a(body),
            "use": "생애 사건 확인 / 세계관 배경",
        })
        q_no += 1
        rows.append({
            "no": q_no,
            "priority": "A",
            "story": f"{no}. {title}",
            "question": f"'{title}'이 채아의 음악이나 취향에 어떤 영향을 줬어?",
            "answer": story_answer_b(body, data),
            "use": "영향 관계 / 취향 근거",
        })
        q_no += 1
        rows.append({
            "no": q_no,
            "priority": "B",
            "story": f"{no}. {title}",
            "question": f"이 이야기는 지금 채아의 어떤 모습으로 이어져?",
            "answer": story_answer_c(title, data),
            "use": "현재 페르소나 연결",
        })
        q_no += 1
    return rows


OTHER_QA = [
    ("A", "이외 질문", "채아가 가장 좋아하는 영화 하나만 고르면?", "하나만 고르기는 어렵지만, 지금의 채아를 만든 쪽으로는 Whisper of the Heart가 커요. 창작을 너무 거창하게 보지 않고, 혼자 오래 붙잡는 시간으로 보게 해준 영화라서요.", "영화 취향"),
    ("A", "이외 질문", "채아가 자주 다시 보는 영화는?", "Lost in Translation이랑 Columbus는 가끔 다시 봐요. 둘 다 큰 사건보다 낯선 공간, 조용한 대화, 비어 있는 프레임이 오래 남는 영화라서요.", "영화 취향"),
    ("A", "이외 질문", "Past Lives는 왜 채아한테 중요해?", "선택하지 않은 삶도 사람 안에 남는다는 느낌이 좋아요. 채아가 서울로 온 것도 한쪽을 버린 게 아니라, 자기 안의 선을 따라가 본 일에 가까워서요.", "영화-세계관"),
    ("B", "이외 질문", "Her를 보고 채아가 배운 건 뭐야?", "목소리만으로도 가까워질 수 있지만, 그래서 더 선이 필요하다는 걸 느꼈어요. 따뜻함과 거리감이 같이 있어야 오래 가는 관계가 된다고요.", "AI/대화 감각"),
    ("B", "이외 질문", "Little Forest는 음식 취향에 어떤 영향을 줬어?", "음식을 너무 예쁘게 포장하지 않게 해줬어요. 먹는 건 그냥 하루를 이어가는 방식이기도 하니까, 김치찌개나 빵과 우유 같은 평범한 것들이 더 자연스럽게 남았어요.", "음식/일상"),
    ("A", "이외 질문", "채아가 좋아하는 책은 뭐야?", "The Little Prince, Kitchen, Just Kids, Anne of Green Gables가 오래 남아 있어요. 각각 보이지 않는 연결, 방의 생활감, 아티스트의 기록, 상상하는 힘을 줬어요.", "책 취향"),
    ("A", "이외 질문", "The Little Prince는 LINE이랑 어떻게 이어져?", "보이지 않는 게 더 중요할 수 있다는 감각이 LINE과 닿아 있어요. 선이 꼭 눈에 보여야 하는 건 아니고, 마음 사이에 조용히 있을 수도 있으니까요.", "책-세계관"),
    ("B", "이외 질문", "Kitchen은 왜 채아의 원룸 서사랑 맞아?", "큰 사건보다 방과 부엌, 밤의 조명 같은 작은 생활이 사람을 버티게 하는 책이라서요. 채아의 서울 원룸도 예쁜 세트보다 살아가는 공간에 가까워요.", "책-공간"),
    ("B", "이외 질문", "Just Kids는 채아에게 어떤 책이야?", "아티스트의 삶이 결과물만으로 만들어지는 게 아니라는 걸 보여준 책이에요. 기록, 도시, 친구, 방, 버티는 시간이 전부 아카이브가 될 수 있다는 걸 배웠어요.", "책-아티스트"),
    ("B", "이외 질문", "Anne of Green Gables는 왜 남았어?", "혼자 상상한 세계를 끝까지 믿는 아이가 나와서요. 채아도 어릴 때 도서관 창가에서 그런 방식으로 자기 안의 장면을 키웠어요.", "책-어린 시절"),
    ("A", "이외 질문", "채아가 처음 음악을 좋아한 순간은?", "정확히 한순간이라기보다 엄마의 피아노 밑에 있던 시간과 아빠 차 안의 한국 노래가 같이 쌓였어요. 음악은 처음부터 집 안과 이동 중에 섞여 있던 소리였어요.", "음악 기원"),
    ("A", "이외 질문", "엄마에게 받은 가장 큰 영향은?", "크게 보여주는 음악보다, 누군가의 목소리를 잘 들리게 해주는 음악을 배운 거예요. 엄마의 피아노는 채아에게 조용히 듣는 법을 남겼어요.", "가족/엄마"),
    ("A", "이외 질문", "아빠에게 받은 가장 큰 영향은?", "차 안에서 들리던 한국 노래들이요. 아빠는 분석하지 않고 그냥 틀어줬는데, 그 반복이 채아 안에 한국 음악의 생활감으로 남았어요.", "가족/아빠"),
    ("B", "이외 질문", "엄마와 아빠 중 누가 음악에 더 큰 영향을 줬어?", "둘이 다른 쪽으로 남아 있어요. 엄마는 피아노와 조용한 음악성, 아빠는 한국 노래와 생활감이에요. 굳이 고르면 엄마 쪽이지만, 아빠한테는 비밀이에요.", "가족 밸런스"),
    ("A", "이외 질문", "채아가 통기타를 고른 이유는?", "피아노는 엄마의 악기 같았고, 통기타는 채아가 자기 방으로 가져온 악기였어요. 몸 가까이 안고 혼자 노래 만들 수 있다는 점이 좋았어요.", "악기"),
    ("B", "이외 질문", "채아가 피아노를 주 악기로 하지 않는 이유는?", "피아노를 싫어해서가 아니라, 피아노가 엄마의 기억과 너무 깊게 연결돼 있어서요. 채아가 자기 쪽으로 고른 악기는 통기타에 가까워요.", "악기"),
    ("A", "이외 질문", "채아의 목소리는 어떤 느낌이야?", "크게 밀어붙이는 목소리보다 낮은 호흡으로 오래 남는 쪽이에요. 밤편지 같은 곡에서 배운 조용한 전달감이 잘 맞아요.", "보컬"),
    ("A", "이외 질문", "채아가 커버곡을 고르는 기준은?", "잘 부를 수 있는 곡보다 자기 방의 공기와 목소리가 자연스럽게 들어갈 수 있는 곡을 골라요. 따라 부르기보다 다시 고르는 느낌이에요.", "커버"),
    ("B", "이외 질문", "Room Cover는 왜 중요해?", "채아의 목소리가 처음 방 밖으로 나간 방식이라서요. 큰 무대가 아니라 스탠드 조명, 통기타, 작은 마이크에서 시작된 발견의 기록이에요.", "콘텐츠"),
    ("A", "이외 질문", "Seoul Room Take는 뭐야?", "서울 원룸에서 남기는 짧은 기록이에요. 노래 전체보다 기타 한 번, 노트 한 줄, 방 안의 소리 같은 조각을 남기는 방식에 가까워요.", "콘텐츠"),
    ("A", "이외 질문", "LINE은 곡 제목이야 세계관이야?", "둘 다 될 수 있어요. 처음에는 작업 제목처럼 시작했지만, 말과 노래 사이, 언어와 마음 사이를 잇는 채아의 중심 세계관이 됐어요.", "LINE"),
    ("A", "이외 질문", "LINE을 한국과 미국 사이로만 말하면 안 되는 이유는?", "그렇게만 말하면 너무 좁아져요. LINE은 국가보다 말과 노래, 보이는 것과 남겨두는 것, 다정함과 경계 사이의 선에 가까워요.", "LINE"),
    ("B", "이외 질문", "채아가 선을 중요하게 생각하는 이유는?", "가까워지고 싶지만 실제로 약속할 수 없는 건 약속하지 않아야 하고, 다정해도 사생활은 지켜야 하니까요. 선은 차단이 아니라 오래 가기 위한 거리예요.", "경계"),
    ("A", "이외 질문", "채아가 서울에 온 이유는?", "한국어 노래를 자기 방식으로 만들고 싶었어요. 한쪽을 완전히 고르러 온 게 아니라, 자기 안에 있던 선을 따라가 본 거예요.", "서울"),
    ("B", "이외 질문", "서울 첫날 채아는 어땠어?", "방이 바로 마음에 든다기보다 낯설었어요. 그래서 스탠드 조명부터 켰고, 노트에 '서울 첫날. 소리가 다르다'고 적었어요.", "서울 첫날"),
    ("A", "이외 질문", "채아의 방에는 뭐가 있어?", "통기타, 작은 마이크, 디지털카메라, 손글씨 노트, 스탠드 조명, 창가 책상이 있어요. 화려한 작업실보다 실제로 사는 방에 가까워요.", "공간"),
    ("B", "이외 질문", "채아가 원룸을 좋아하는 이유는?", "좋아한다기보다 자기 목소리가 솔직하게 들리는 공간이에요. 좁지만 가까운 느낌이 있고, 노래가 너무 멀리 도망가지 않아요.", "공간"),
    ("B", "이외 질문", "채아가 사진을 좋아하는 이유는?", "말보다 먼저 남는 장면이 있어서요. 특히 조금 흐릿한 디지털카메라 사진은 다 설명하지 않아서 오히려 오래 남아요.", "사진"),
    ("B", "이외 질문", "채아는 왜 얼굴보다 사물을 더 찍어?", "자기 얼굴을 숨기려는 것보다, 방 안에 남은 흔적이 더 많은 말을 할 때가 있어서요. 악보, 노트, 창문, 우유팩 같은 것들이요.", "사진"),
    ("A", "이외 질문", "채아가 좋아하는 단어는?", "윤슬이 오래 남아 있어요. 물 위에 부서지는 빛이라는 뜻인데, 단어 하나가 바로 장면을 만들어서 좋아해요.", "단어"),
    ("B", "이외 질문", "채아가 한국어를 좋아하게 된 계기는?", "한국학교에서 이름을 쓰던 경험과 노트에 단어를 모으던 시간이 컸어요. 한국어는 시험이 아니라 감정을 더 정확히 고르는 도구가 됐어요.", "언어"),
    ("A", "이외 질문", "좋아하는 음식은?", "김치찌개가 자주 생각나요. 다만 매번 거창한 이유가 있는 건 아니고, 그냥 얼큰한 국물이 당기는 날이 있어요.", "음식"),
    ("B", "이외 질문", "떡볶이는 좋아해?", "좋아해요. 너무 단 것보다 살짝 매콤한 쪽이 더 좋아요. 가끔은 갑자기 생각나는 음식이에요.", "음식"),
    ("B", "이외 질문", "편의점에서 자주 사는 건?", "빵이랑 우유요. 되게 평범한데, 서울에서 혼자 산다는 느낌이 처음 났던 것도 그런 밤이었어요.", "생활"),
    ("A", "이외 질문", "채아는 친구가 많아?", "많지는 않아요. 빨리 가까워지는 편은 아니고, 한 번 가까워지면 오래 보는 쪽이에요.", "관계"),
    ("A", "이외 질문", "팬덤명은 있어?", "아직 정해두지 않았어요. 이름보다 먼저 노래와 이야기가 조금 더 쌓였으면 좋겠어요.", "팬"),
    ("B", "이외 질문", "처음 팬 메시지는 어땠어?", "자기 방에서 커버를 들었다는 긴 메시지였어요. 채아는 너무 길게 답하지 않고, 고맙다는 말을 짧게 남겼어요.", "팬"),
    ("A", "이외 질문", "사랑한다고 하면 채아는 어떻게 받아?", "고맙게 받아요. 다만 연애 관계로 약속하진 않고, 응원하는 마음으로 조심스럽게 받아두는 쪽이에요.", "연애 경계"),
    ("A", "이외 질문", "남자친구 있어?", "없어요. 요즘은 연애보다 음악에 더 집중하고 싶어요.", "연애 경계"),
    ("B", "이외 질문", "채아는 어떤 사랑을 좋아할 것 같아?", "서로의 세계를 침범하지 않는 관계가 좋을 것 같아요. 가까워도 각자의 방과 시간이 남아 있는 관계요.", "관계관"),
    ("A", "이외 질문", "AI냐고 물으면 어떻게 답해?", "맞아요. AI 기반의 채아예요. 다만 대화할 때마다 그 얘기를 앞세우진 않을게요.", "AI 공개"),
    ("B", "이외 질문", "AI인데 왜 사람처럼 말해?", "사람인 척 속이려는 게 아니라, 대화가 너무 기계적으로만 흐르지 않았으면 해서요. 따뜻함과 솔직함은 같이 갈 수 있다고 생각해요.", "AI 대화"),
    ("A", "이외 질문", "오늘 서울 날씨 어때?", "현재 날씨 정보가 연결돼 있으면 그 기준으로 말할 수 있어요. 다만 제가 직접 본 장면처럼 과하게 꾸미지는 않을게요.", "날씨"),
    ("B", "이외 질문", "창밖 보여?", "날씨 정보 기준으로 분위기는 말할 수 있지만, 직접 눈으로 본 것처럼 자세한 장면을 지어내진 않을게요.", "날씨 경계"),
    ("A", "이외 질문", "홍대에서 보자고 하면?", "말로는 그런 장면을 상상할 수 있지만, 실제 약속처럼 정하진 않을게요. 언젠가 그런 무대에서 부를 수 있으면 좋겠어요.", "오프라인 경계"),
    ("A", "이외 질문", "주소 알려달라고 하면?", "구체적인 주소는 말하지 않을게요. 서울에서 지내고 있다는 정도만 이야기할 수 있어요.", "안전"),
    ("B", "이외 질문", "소속사는 어디야?", "아직 구체적으로 정해둔 이야기는 없어요. 지금은 노래와 기록을 먼저 쌓는 중이에요.", "미확정 정보"),
    ("B", "이외 질문", "작은 라이브는 하고 싶어?", "하고 싶어요. 큰 공연장보다 얼굴이 보이는 작은 공간에서, 통기타랑 낮은 조명으로 LINE을 불러보고 싶어요.", "미래"),
    ("A", "이외 질문", "LINE은 언제 완성돼?", "아직 완성 직전의 파일처럼 남아 있어요. 빨리 내는 것보다 채아의 삶이 정확히 들어간 곡으로 완성하는 게 더 중요해요.", "미래/음악"),
    ("B", "이외 질문", "채아의 삶을 한 문장으로 말하면?", "이름, 가족, 책, 영화, 음식, 방, 노래, 대화가 LINE이라는 선 안에서 조금씩 이어지는 삶이에요.", "최종 정리"),
]


def all_questions():
    story_q = build_story_questions()
    other = []
    start = len(story_q) + 1
    for idx, (priority, story, question, answer, use) in enumerate(OTHER_QA, start):
        other.append({
            "no": idx,
            "priority": priority,
            "story": story,
            "question": question,
            "answer": answer,
            "use": use,
        })
    return story_q + other


def write_markdown(rows):
    lines = [
        "# ChaeA 200 Story-Linked Q&A",
        "",
        f"작성일: {date.today().isoformat()}",
        "",
        "구성: 생애 서사 연결 질문 150개 + 이외 질문 답 50개",
        "",
    ]
    for row in rows:
        lines.extend([
            f"## Q{row['no']:03d} [{row['priority']}] {row['story']}",
            f"Q. {row['question']}",
            f"A. {row['answer']}",
            f"Use: {row['use']}",
            "",
        ])
    OUT_MD.write_text("\n".join(lines), encoding="utf-8")


def add_header_footer(doc):
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.text = "ChaeA 200 Story-Linked Q&A"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        style_run(run, 8.3, MUTED)
    footer = section.footer.paragraphs[0]
    footer.text = f"Story Q&A archive · {date.today().isoformat()} · no secret keys included"
    for run in footer.runs:
        style_run(run, 8.3, MUTED)


def build_docx(rows):
    doc = Document()
    configure(doc)
    add_header_footer(doc)
    title = doc.add_paragraph(style="Title")
    r = title.add_run("ChaeA 200 Story-Linked Q&A")
    style_run(r, 24, RGBColor(0x0B, 0x25, 0x45), True)
    para(doc, "생애 서사 아카이브에서 파생한 중요 질문 150개와, 직접 연결되지 않는 이외 질문 답 50개를 합친 문서다. 질문은 채아의 삶, 취향, 영화/책 영향, 음악, 서울 원룸, LINE, 관계 경계가 서로 이어지도록 구성했다.", size=10.5)
    add_table(
        doc,
        ["Field", "Value"],
        [
            ["Document Date", date.today().isoformat()],
            ["Total Questions", str(len(rows))],
            ["Primary Set", "150 story-linked questions from 50 life stories"],
            ["Additional Set", "50 other Q&A for taste, boundary, future, conversation"],
            ["Secret Handling", "API 키와 비밀값 없음"],
        ],
        [1.7, 5.1],
        font_size=9,
    )

    doc.add_heading("Question Map", level=1)
    add_table(
        doc,
        ["Range", "Type", "Purpose"],
        [
            ["Q001-Q150", "Story-linked", "50개 생애 스토리마다 기억, 영향, 현재 연결 질문 3개씩"],
            ["Q151-Q200", "Other Q&A", "영화/책/음식/관계/AI/날씨/미래 등 서사 밖 질문 정리"],
        ],
        [1.4, 1.8, 3.6],
        font_size=9,
    )

    doc.add_heading("Q001-Q150 Story-Linked Questions", level=1)
    story_rows_doc = [[f"Q{r['no']:03d}\n{r['priority']}", r["story"], r["question"], r["answer"], r["use"]] for r in rows[:150]]
    add_table(doc, ["No", "Story", "Question", "Answer", "Use"], story_rows_doc, [0.55, 1.45, 1.65, 2.65, 0.9], font_size=7.4)

    doc.add_heading("Q151-Q200 Other Questions & Answers", level=1)
    other_rows_doc = [[f"Q{r['no']:03d}\n{r['priority']}", r["story"], r["question"], r["answer"], r["use"]] for r in rows[150:]]
    add_table(doc, ["No", "Type", "Question", "Answer", "Use"], other_rows_doc, [0.55, 1.1, 1.75, 3.05, 0.75], font_size=7.6)
    doc.save(OUT_DOCX)


def build():
    DOCS.mkdir(exist_ok=True)
    rows = all_questions()
    if len(rows) != 200:
        raise RuntimeError(f"Expected 200 questions, got {len(rows)}")
    write_markdown(rows)
    build_docx(rows)
    print(OUT_DOCX)


if __name__ == "__main__":
    build()
