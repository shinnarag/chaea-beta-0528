from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]

PRESET = {
    "page_width": Inches(8.5),
    "page_height": Inches(11),
    "margin": Inches(1),
    "base_font": "Calibri",
    "east_asia_font": "Apple SD Gothic Neo",
    "body_size": Pt(11),
    "body_after": Pt(6),
    "body_line": 1.25,
    "h1_size": Pt(16),
    "h2_size": Pt(13),
    "h3_size": Pt(12),
    "h1_before": Pt(18),
    "h1_after": Pt(10),
    "h2_before": Pt(14),
    "h2_after": Pt(7),
    "h3_before": Pt(10),
    "h3_after": Pt(5),
    "heading_blue": RGBColor(0x2E, 0x74, 0xB5),
    "heading_dark": RGBColor(0x1F, 0x4D, 0x78),
    "muted": RGBColor(0x66, 0x66, 0x66),
    "table_header_fill": "E8EEF5",
    "table_border": "C8D2DE",
}


DOCUMENTS = [
    {
        "source": ROOT / "docs/current/ChaeA_Current_Persona_Master.md",
        "output": ROOT / "docs/current/ChaeA_Current_Persona_Master.docx",
        "title": "ChaeA Current Persona Master",
        "subtitle": "Latest ChaeA persona master for local operations",
    },
    {
        "source": ROOT / "docs/current/ChaeA_Current_Story_Persona_Final.md",
        "output": ROOT / "docs/current/ChaeA_Current_Story_Persona_Final.docx",
        "title": "ChaeA Current Story Persona Final",
        "subtitle": "Final story persona promoted from archived ChaeA lore",
    },
]

REFERENCE_IMAGES = [
    {
        "path": ROOT / "assets/reference/chaea/ChaeA-1.jpg",
        "caption": "ChaeA portrait reference: long wavy hair, right-eye mole, soft peach makeup",
    },
    {
        "path": ROOT / "assets/reference/chaea/ChaeA-3.jpg",
        "caption": "Outdoor daily reference: natural profile, brown highlights, casual denim mood",
    },
    {
        "path": ROOT / "assets/reference/chaea/ChaeA-10.png",
        "caption": "Room portrait reference: Seoul room, bed/window light, relaxed home styling",
    },
    {
        "path": ROOT / "assets/reference/cloth/cloth_10.jpg",
        "caption": "Clothing reference: pale blue stripe top and loose ivory pants",
    },
    {
        "path": ROOT / "assets/reference/place/chaea room_1.png",
        "caption": "Room reference: warm one-room layout, bed, curtain, lamp, small storage",
    },
    {
        "path": ROOT / "assets/reference/place/chaea room_curtain.png",
        "caption": "Room reference: evening curtain mood, muted beige-gray palette",
    },
]


def main() -> None:
    for item in DOCUMENTS:
        build_docx(item)
        print(f"wrote {item['output'].relative_to(ROOT)}")


def build_docx(item: dict[str, Path | str]) -> None:
    source = Path(item["source"])
    output = Path(item["output"])
    frontmatter, body = split_frontmatter(source.read_text(encoding="utf-8"))

    doc = Document()
    configure_document(doc, str(item["title"]))
    add_title_block(doc, str(item["title"]), str(item["subtitle"]), source)
    if frontmatter:
        add_metadata_table(doc, frontmatter)
    parse_markdown_body(doc, body, str(item["title"]))
    add_visual_reference_appendix(doc)
    doc.save(output)


def configure_document(doc: Document, running_title: str) -> None:
    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = PRESET["page_width"]
    section.page_height = PRESET["page_height"]
    section.top_margin = PRESET["margin"]
    section.bottom_margin = PRESET["margin"]
    section.left_margin = PRESET["margin"]
    section.right_margin = PRESET["margin"]
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    configure_style(styles["Normal"], PRESET["base_font"], PRESET["body_size"], RGBColor(0, 0, 0))
    styles["Normal"].paragraph_format.space_after = PRESET["body_after"]
    styles["Normal"].paragraph_format.line_spacing = PRESET["body_line"]

    for style_name, size, color, before, after in [
        ("Heading 1", PRESET["h1_size"], PRESET["heading_blue"], PRESET["h1_before"], PRESET["h1_after"]),
        ("Heading 2", PRESET["h2_size"], PRESET["heading_blue"], PRESET["h2_before"], PRESET["h2_after"]),
        ("Heading 3", PRESET["h3_size"], PRESET["heading_dark"], PRESET["h3_before"], PRESET["h3_after"]),
    ]:
        style = styles[style_name]
        configure_style(style, PRESET["base_font"], size, color, bold=True)
        style.paragraph_format.space_before = before
        style.paragraph_format.space_after = after
        style.paragraph_format.keep_with_next = True

    for style_name in ["List Bullet", "List Number"]:
        style = styles[style_name]
        configure_style(style, PRESET["base_font"], PRESET["body_size"], RGBColor(0, 0, 0))
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = PRESET["body_line"]

    header = section.header.paragraphs[0]
    header.text = running_title
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = PRESET["muted"]
    set_run_font(header.runs[0])

    footer = section.footer.paragraphs[0]
    footer.text = "ChaeA local persona archive"
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = PRESET["muted"]
    set_run_font(footer.runs[0])


def configure_style(style, font_name, size, color, bold=False) -> None:
    style.font.name = font_name
    style.font.size = size
    style.font.color.rgb = color
    style.font.bold = bold
    style._element.rPr.rFonts.set(qn("w:eastAsia"), PRESET["east_asia_font"])


def add_title_block(doc: Document, title: str, subtitle: str, source: Path) -> None:
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_after = Pt(3)
    title_run = title_p.add_run(title)
    title_run.font.name = PRESET["base_font"]
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)
    set_run_font(title_run)

    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_after = Pt(12)
    subtitle_run = subtitle_p.add_run(subtitle)
    subtitle_run.font.size = Pt(10)
    subtitle_run.font.color.rgb = PRESET["muted"]
    set_run_font(subtitle_run)

    source_p = doc.add_paragraph()
    source_p.paragraph_format.space_after = Pt(12)
    run = source_p.add_run(f"Source: {source.relative_to(ROOT)}")
    run.font.size = Pt(9)
    run.font.color.rgb = PRESET["muted"]
    set_run_font(run)


def add_metadata_table(doc: Document, frontmatter: str) -> None:
    heading = doc.add_paragraph("Document Metadata", style="Heading 1")
    heading.paragraph_format.keep_with_next = True
    rows = parse_frontmatter(frontmatter)
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    set_table_geometry(table, [2700, 6660])
    set_table_borders(table)
    hdr = table.rows[0].cells
    hdr[0].text = "Field"
    hdr[1].text = "Value"
    for cell in hdr:
        format_header_cell(cell)
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value
        for cell in cells:
            format_body_cell(cell)


def parse_markdown_body(doc: Document, body: str, doc_title: str) -> None:
    lines = body.splitlines()
    paragraph_buffer: list[str] = []
    index = 0
    first_h1_skipped = False

    def flush_paragraph() -> None:
        if not paragraph_buffer:
            return
        text = " ".join(part.strip() for part in paragraph_buffer if part.strip())
        paragraph_buffer.clear()
        if text:
            add_markdown_paragraph(doc, text)

    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()

        if not stripped:
            flush_paragraph()
            index += 1
            continue

        if is_table_start(lines, index):
            flush_paragraph()
            table_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            add_markdown_table(doc, table_lines)
            continue

        heading_match = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading_match:
            flush_paragraph()
            level = len(heading_match.group(1))
            text = strip_inline_markdown(heading_match.group(2))
            if level == 1 and not first_h1_skipped and normalize_title(text) == normalize_title(doc_title):
                first_h1_skipped = True
                index += 1
                continue
            style = "Heading 1" if level <= 2 else "Heading 2" if level == 3 else "Heading 3"
            doc.add_paragraph(text, style=style)
            index += 1
            continue

        if stripped.startswith(">"):
            flush_paragraph()
            quote = stripped.lstrip(">").strip()
            add_quote(doc, quote)
            index += 1
            continue

        bullet_match = re.match(r"^-\s+(.*)$", stripped)
        if bullet_match:
            flush_paragraph()
            add_markdown_paragraph(doc, bullet_match.group(1), style="List Bullet")
            index += 1
            continue

        number_match = re.match(r"^\d+\.\s+(.*)$", stripped)
        if number_match:
            flush_paragraph()
            add_markdown_paragraph(doc, number_match.group(1), style="List Number")
            index += 1
            continue

        paragraph_buffer.append(stripped)
        index += 1

    flush_paragraph()


def add_markdown_paragraph(doc: Document, text: str, style: str | None = None):
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(4 if style else 6)
    paragraph.paragraph_format.line_spacing = PRESET["body_line"]
    add_runs_from_markdown(paragraph, text)
    return paragraph


def add_quote(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.right_indent = Inches(0.15)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(strip_inline_markdown(text))
    run.italic = True
    run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    set_run_font(run)


def add_markdown_table(doc: Document, table_lines: list[str]) -> None:
    rows: list[list[str]] = []
    for line in table_lines:
        cells = [cell.strip() for cell in line.strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
            continue
        rows.append(cells)

    if not rows:
        return

    col_count = max(len(row) for row in rows)
    for row in rows:
        row.extend([""] * (col_count - len(row)))

    table = doc.add_table(rows=1, cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    widths = table_widths(col_count)
    set_table_geometry(table, widths)
    set_table_borders(table)

    for col, text in enumerate(rows[0]):
        table.rows[0].cells[col].text = strip_inline_markdown(text)
        format_header_cell(table.rows[0].cells[col])

    for row in rows[1:]:
        cells = table.add_row().cells
        for col, text in enumerate(row):
            cells[col].text = strip_inline_markdown(text)
            format_body_cell(cells[col])

    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_visual_reference_appendix(doc: Document) -> None:
    doc.add_page_break()
    doc.add_paragraph("Visual Reference Appendix", style="Heading 1")
    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(10)
    add_runs_from_markdown(
        intro,
        "대표 레퍼런스 사진을 문서 안에 임베드했다. 최신 외형 기준은 오른쪽 눈 아래 작은 점, 긴 웨이브 헤어, 다크 브라운과 라이트 브라운 포인트, 얇은 실버 목걸이, 밝은 톤 의상이다.",
    )

    images = [item for item in REFERENCE_IMAGES if Path(item["path"]).exists()]
    if not images:
        add_markdown_paragraph(doc, "No local reference images found.")
        return

    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    set_table_geometry(table, [4680, 4680])
    set_table_borders(table)

    for idx in range(0, len(images), 2):
        cells = table.add_row().cells
        for col in range(2):
            image_idx = idx + col
            cell = cells[col]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell, top=120, bottom=120, start=120, end=120)
            clear_cell(cell)
            if image_idx >= len(images):
                continue
            item = images[image_idx]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            run.add_picture(str(item["path"]), width=Inches(2.8))
            caption = cell.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.space_before = Pt(4)
            caption.paragraph_format.space_after = Pt(0)
            caption_run = caption.add_run(item["caption"])
            caption_run.font.size = Pt(8)
            caption_run.font.color.rgb = PRESET["muted"]
            set_run_font(caption_run)


def table_widths(col_count: int) -> list[int]:
    if col_count == 2:
        return [2700, 6660]
    if col_count == 3:
        return [1700, 3600, 4060]
    if col_count == 4:
        return [1600, 2200, 2760, 2800]
    base = 9360 // col_count
    widths = [base] * col_count
    widths[-1] += 9360 - sum(widths)
    return widths


def format_header_cell(cell) -> None:
    set_cell_shading(cell, PRESET["table_header_fill"])
    format_cell_paragraphs(cell, bold=True, color=RGBColor(0x0B, 0x25, 0x45))


def format_body_cell(cell) -> None:
    format_cell_paragraphs(cell, bold=False, color=RGBColor(0, 0, 0))


def format_cell_paragraphs(cell, bold: bool, color: RGBColor) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell, top=80, bottom=80, start=120, end=120)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.15
        for run in paragraph.runs:
            run.font.name = PRESET["base_font"]
            run.font.size = Pt(9)
            run.font.bold = bold
            run.font.color.rgb = color
            set_run_font(run)


def clear_cell(cell) -> None:
    for paragraph in cell.paragraphs:
        p = paragraph._element
        p.getparent().remove(p)
    cell.add_paragraph()


def add_runs_from_markdown(paragraph, text: str) -> None:
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            add_text_run(paragraph, strip_links(text[pos : match.start()]))
        token = match.group(0)
        if token.startswith("**"):
            add_text_run(paragraph, strip_links(token[2:-2]), bold=True)
        elif token.startswith("`"):
            add_text_run(paragraph, token[1:-1], mono=True)
        pos = match.end()
    if pos < len(text):
        add_text_run(paragraph, strip_links(text[pos:]))


def add_text_run(paragraph, text: str, bold: bool = False, mono: bool = False) -> None:
    if not text:
        return
    run = paragraph.add_run(text)
    run.font.name = "Consolas" if mono else PRESET["base_font"]
    run.font.size = Pt(10 if mono else 11)
    run.font.bold = bold
    if mono:
        run.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)
    set_run_font(run, "Consolas" if mono else PRESET["east_asia_font"])


def split_frontmatter(text: str) -> tuple[str, str]:
    if not text.startswith("---\n"):
        return "", text
    parts = text.split("---", 2)
    if len(parts) < 3:
        return "", text
    return parts[1].strip(), parts[2].lstrip("\n")


def parse_frontmatter(frontmatter: str) -> list[tuple[str, str]]:
    rows: list[tuple[str, str]] = []
    current_key: str | None = None
    current_values: list[str] = []

    def flush() -> None:
        nonlocal current_key, current_values
        if current_key is not None:
            rows.append((current_key, "\n".join(current_values).strip()))
        current_key = None
        current_values = []

    for raw in frontmatter.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        key_match = re.match(r"^([A-Za-z0-9_\-]+):\s*(.*)$", line)
        if key_match:
            flush()
            current_key = key_match.group(1)
            value = key_match.group(2).strip().strip('"')
            current_values = [value] if value else []
            continue
        item_match = re.match(r"^\s*-\s+(.*)$", line)
        if item_match and current_key:
            current_values.append(item_match.group(1).strip())
    flush()
    return rows


def is_table_start(lines: list[str], index: int) -> bool:
    if index + 1 >= len(lines):
        return False
    return lines[index].strip().startswith("|") and re.match(r"^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?$", lines[index + 1].strip()) is not None


def strip_inline_markdown(text: str) -> str:
    text = strip_links(text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return text


def strip_links(text: str) -> str:
    return re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)


def normalize_title(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", text.lower())


def set_run_font(run, east_asia: str | None = None) -> None:
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), east_asia or PRESET["east_asia_font"])


def set_table_geometry(table, widths: list[int]) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])


def set_cell_width(cell, width: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), PRESET["table_border"])


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int, bottom: int, start: int, end: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in [("top", top), ("bottom", bottom), ("start", start), ("end", end)]:
        element = margins.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


if __name__ == "__main__":
    main()
